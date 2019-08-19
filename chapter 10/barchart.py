from openpyxl import load_workbook
wb = load_workbook('example.xlsx')
ws = wb.active
from openpyxl.chart import BarChart, Reference
values = Reference(ws, min_col=2, min_row=2, max_col=4, max_row=7)
ctg=Reference(ws, min_col=1,min_row=3, max_col=1, max_row=7)

c1 = BarChart()
c1.add_data(values, titles_from_data=True
c1.title = "Bar Chart"
c1.x_axis.title = 'Months'
c1.y_axis.title = 'Sales'
ws.add_chart(c1, "A10")
c1.set_categories(ctg)
wb.save(filename='example.xlsx')

from openpyxl.chart import LineChart
c2 = LineChart()
c2.add_data(values, titles_from_data=True)#legends
c2.title = "Line Chart"
c2.x_axis.title = 'Months'
c2.y_axis.title = 'Sales'
ws.add_chart(c2, "F2")
c2.set_categories(ctg)


